"""Helpers to build realistic test workbooks (formatting, filters, freeze panes, extra sheets)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side


def make_workbook(
    path: Path,
    headers: list[str],
    rows: list[list[Any]],
    *,
    title: str = "Streamers",
    extra_sheet: bool = True,
    header_offset: int = 0,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for _ in range(header_offset):
        ws.append(["Streamer list — Q3"])
    ws.append(headers)
    hdr = header_offset + 1
    thin = Side(style="thin", color="999999")
    for cell in ws[hdr]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.border = Border(bottom=thin)
    for r in rows:
        ws.append(r)
    for letter, width in zip("ABCDEFG", (22, 14, 24, 28, 12, 12, 12), strict=False):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{hdr}:{chr(64 + len(headers))}{hdr + len(rows)}"
    ws.cell(row=hdr + 1, column=2).number_format = "@"
    if extra_sheet:
        other = wb.create_sheet("Notes")
        other["A1"] = "Do not modify"
        other["B2"] = "=1+1"
        other.merge_cells("A4:C4")
        other["A4"] = "merged"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


KICK_HEADERS = ["id_kick", "country", "id_twitch", "remarks"]
TWITCH_HEADERS = ["id_twitch", "country", "id_kick", "remarks"]
