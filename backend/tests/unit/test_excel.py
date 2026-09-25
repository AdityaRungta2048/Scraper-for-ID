"""Importer / exporter / verifier tests (no network)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.excel.exporter import CellWrite, ExportError, export_processed, output_filename
from app.excel.importer import WorkbookValidationError, analyze_workbook, normalize_header, read_rows
from app.excel.review_report import review_filename
from app.excel.verifier import verify_output
from tests.excel_helpers import KICK_HEADERS, TWITCH_HEADERS, make_workbook


def kick_book(tmp_path: Path, rows=None, **kw) -> Path:
    rows = rows or [
        ["trsniIs", "Germany", None, None],
        ["tittoshow", "Spain", None, None],
        ["nikkilve", "Germany", None, None],
        ["teufeurs", "France", None, None],
    ]
    return make_workbook(tmp_path / "streamers.xlsx", KICK_HEADERS, rows, **kw)


def test_header_normalisation():
    assert normalize_header("  ID_Kick ") == "id_kick"
    assert normalize_header("id twitch") == "id_twitch"
    assert normalize_header("Remarks") == "remarks"
    assert normalize_header("Kick-ID") == "id_kick"


def test_detect_kick_source(tmp_path):
    a = analyze_workbook(kick_book(tmp_path))
    assert a.detected_platform == "kick" and not a.ambiguous
    assert (a.columns.id_kick, a.columns.country, a.columns.id_twitch, a.columns.remarks) == (1, 2, 3, 4)
    assert a.data_rows == 4


def test_detect_twitch_source_with_none_placeholders(tmp_path):
    rows = [
        ["bobinice7676", "France", "bobinice76", None],
        ["kosstochka", "Italy", "None", None],
        ["primeleague", "Germany", "None", None],
        ["starwraith", "Spain", "starwraith", None],
    ]
    a = analyze_workbook(make_workbook(tmp_path / "t.xlsx", TWITCH_HEADERS, rows))
    assert a.detected_platform == "twitch" and not a.ambiguous


def test_ambiguous_when_order_and_fill_disagree(tmp_path):
    rows = [[None, "France", "abc", None], [None, "Spain", "def", None], [None, "Italy", "ghi", None]]
    a = analyze_workbook(make_workbook(tmp_path / "amb.xlsx", KICK_HEADERS, rows))
    assert a.ambiguous and a.detected_platform is None


def test_header_row_offset_and_case_variation(tmp_path):
    p = make_workbook(
        tmp_path / "o.xlsx",
        ["ID_Kick", " Country ", "ID Twitch", "Remarks"],
        [["abc", "France", None, None]],
        header_offset=2,
    )
    a = analyze_workbook(p)
    assert a.header_row == 3 and a.detected_platform == "kick"


def test_missing_headers_rejected(tmp_path):
    p = make_workbook(tmp_path / "bad.xlsx", ["name", "country"], [["x", "y"]])
    with pytest.raises(WorkbookValidationError, match="Required headers"):
        analyze_workbook(p)


def test_missing_remarks_column_is_added(tmp_path):
    p = make_workbook(tmp_path / "nr.xlsx", ["id_kick", "country", "id_twitch"], [["abc", "France", None]])
    a = analyze_workbook(p)
    assert a.columns.remarks is None and a.columns.remarks_new_col == 4
    assert any("remarks" in w for w in a.warnings)


def test_invalid_files(tmp_path):
    bad = tmp_path / "x.xlsx"
    bad.write_bytes(b"this is not a zip")
    with pytest.raises(WorkbookValidationError, match="not a valid Excel"):
        analyze_workbook(bad)
    legacy = tmp_path / "x.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(WorkbookValidationError, match="xls"):
        analyze_workbook(legacy)
    with pytest.raises(WorkbookValidationError, match="Unsupported"):
        analyze_workbook(tmp_path / "x.csv")


def test_no_data_rows(tmp_path):
    with pytest.raises(WorkbookValidationError, match="no data rows"):
        analyze_workbook(make_workbook(tmp_path / "e.xlsx", KICK_HEADERS, []))


def test_read_rows_keeps_original_row_numbers_duplicates_and_blanks(tmp_path):
    rows = [
        ["a", "France", None, None],
        [None, "Spain", None, None],
        ["a", "France", None, None],
        ["b", "Italy", "pre", "note"],
    ]
    p = kick_book(tmp_path, rows)
    a = analyze_workbook(p)
    got = read_rows(p, a.columns, "kick")
    assert [r.original_row for r in got] == [2, 3, 4, 5]
    assert [r.source_value for r in got] == ["a", None, "a", "b"]
    assert got[3].existing_destination == "pre" and got[3].existing_remarks == "note"


def test_output_filename():
    assert output_filename("streamers.xlsx") == "streamers_processed.xlsx"
    assert output_filename("my_streamers.xlsx") == "my_streamers_processed.xlsx"
    assert output_filename("macro.xlsm") == "macro_processed.xlsm"
    assert review_filename("my_streamers.xlsx") == "my_streamers_review.xlsx"


def _export(tmp_path, writes, rows=None):
    src = kick_book(tmp_path, rows)
    a = analyze_workbook(src)
    out = tmp_path / "out" / output_filename(src.name)
    expected = export_processed(src, out, a.columns, "kick", writes)
    return src, out, a, expected


def test_export_writes_only_target_cells_and_preserves_everything(tmp_path):
    writes = [CellWrite(3, True, "tittoshow", False, None), CellWrite(5, False, None, True, "no kick id")]
    src, out, a, expected = _export(tmp_path, writes)
    wb = load_workbook(out)
    ws = wb["Streamers"]
    assert [ws.cell(row=r, column=1).value for r in range(1, 6)] == [
        "id_kick",
        "trsniIs",
        "tittoshow",
        "nikkilve",
        "teufeurs",
    ]
    assert ws.cell(row=3, column=3).value == "tittoshow"
    assert ws.cell(row=5, column=4).value == "no kick id"
    assert ws.cell(row=2, column=3).value is None
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D5"
    assert ws.cell(row=1, column=1).font.bold and ws.column_dimensions["A"].width == 22
    assert wb.sheetnames == ["Streamers", "Notes"] and wb["Notes"]["B2"].value == "=1+1"
    rep = verify_output(
        src,
        out,
        a.columns,
        "kick",
        expected,
        [(2, "trsniIs"), (3, "tittoshow"), (4, "nikkilve"), (5, "teufeurs")],
    )
    assert rep.ok, rep.to_dict()


def test_verifier_detects_shifted_or_tampered_output(tmp_path):
    writes = [CellWrite(3, True, "tittoshow", False, None)]
    src, out, a, expected = _export(tmp_path, writes)
    wb = load_workbook(out)
    ws = wb["Streamers"]
    # simulate an off-by-one bug: value written one row too low, and a source id changed
    ws.cell(row=3, column=3).value = None
    ws.cell(row=4, column=3).value = "tittoshow"
    ws.cell(row=2, column=1).value = "changed"
    wb.save(out)
    rep = verify_output(src, out, a.columns, "kick", expected, [(2, "trsniIs"), (3, "tittoshow")])
    assert not rep.ok
    failed = {c.name for c in rep.checks if not c.ok}
    assert "cell_values_and_formatting[Streamers]" in failed and "row_identity" in failed
    assert any("C4" in m for m in rep.mismatches)


def test_verifier_detects_deleted_row_and_extra_sheet_change(tmp_path):
    src, out, a, expected = _export(tmp_path, [])
    wb = load_workbook(out)
    wb["Streamers"].delete_rows(3)
    wb["Notes"]["A1"] = "edited"
    wb.save(out)
    rep = verify_output(src, out, a.columns, "kick", expected, [(3, "tittoshow")])
    assert not rep.ok
    assert not next(c for c in rep.checks if c.name == "cell_values_and_formatting[Notes]").ok


def test_verifier_detects_corrupt_file(tmp_path):
    src, out, a, expected = _export(tmp_path, [])
    out.write_bytes(b"garbage")
    rep = verify_output(src, out, a.columns, "kick", expected, [])
    assert not rep.ok


def test_exporter_rejects_duplicate_and_header_writes(tmp_path):
    with pytest.raises(ExportError):
        _export(tmp_path, [CellWrite(3, True, "x", False, None), CellWrite(3, True, "y", False, None)])
    with pytest.raises(ExportError):
        _export(tmp_path, [CellWrite(1, True, "x", False, None)])


def test_export_adds_remarks_header_when_missing(tmp_path):
    src = make_workbook(tmp_path / "nr.xlsx", ["id_kick", "country", "id_twitch"], [["abc", "France", None]])
    a = analyze_workbook(src)
    out = tmp_path / "nr_processed.xlsx"
    expected = export_processed(src, out, a.columns, "kick", [CellWrite(2, False, None, True, "no kick id")])
    ws = load_workbook(out).active
    assert ws["D1"].value == "remarks" and ws["D2"].value == "no kick id"
    assert verify_output(src, out, a.columns, "kick", expected, [(2, "abc")]).ok


def test_link_columns_placed_after_all_existing_columns(tmp_path):
    p = make_workbook(tmp_path / "wide.xlsx", [*KICK_HEADERS, "notes"], [["abc", "France", None, None, "x"]])
    a = analyze_workbook(p)
    assert (a.columns.twitch_link, a.columns.kick_link) == (6, 7)
    assert a.columns.new_headers == {"6": "twitch_id_link", "7": "kick_id_link"}


def test_verifier_detects_wrong_hyperlink(tmp_path):
    writes = [CellWrite(2, False, None, False, None, {"twitch": "https://www.twitch.tv/a", "kick": None})]
    src, out, a, expected = _export(tmp_path, writes)
    wb = load_workbook(out)
    wb["Streamers"]["E2"].hyperlink = "https://evil.example.com"
    wb.save(out)
    rep = verify_output(src, out, a.columns, "kick", expected, [(2, "trsniIs")])
    assert not rep.ok and not next(c for c in rep.checks if c.name == "hyperlinks[Streamers]").ok
