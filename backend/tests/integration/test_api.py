"""HTTP API: upload -> detect -> start -> progress -> download (+ review workflow)."""

from __future__ import annotations

import io
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app
from app.workers.queue import wait_inline
from tests.excel_helpers import KICK_HEADERS, TWITCH_HEADERS, make_workbook
from tests.scenarios import build_kick_world

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client(fake):
    with TestClient(create_app()) as c:
        yield c


def upload(client, path: Path, name: str | None = None):
    return client.post("/api/jobs", files={"file": (name or path.name, path.read_bytes(), XLSX)})


def run_to_completion(client, job_id: str, platform: str | None = None) -> dict:
    r = client.post(f"/api/jobs/{job_id}/start", json={"source_platform": platform})
    assert r.status_code == 200, r.text
    wait_inline(job_id, timeout=60)
    for _ in range(100):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("COMPLETED", "PARTIAL", "FAILED", "CANCELLED"):
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_health_and_config_do_not_leak_secrets(client):
    assert client.get("/api/health").json()["status"] == "ok"
    cfg = client.get("/api/config").json()
    assert cfg["twitch_configured"] and cfg["kick_configured"]
    assert "secret" not in str(cfg).lower() or all("secret" not in k for k in cfg)
    assert "test-twitch-secret" not in str(cfg)


def test_full_flow_kick_workbook(client, fake, tmp_path):
    expects = build_kick_world(fake)
    path = make_workbook(
        tmp_path / "my_streamers.xlsx", KICK_HEADERS, [[e.source, e.country, None, None] for e in expects]
    )
    r = upload(client, path)
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["detected_platform"] == "kick" and job["total_rows"] == len(expects)
    assert job["status"] == "UPLOADED" and not job["needs_platform_choice"]

    job = run_to_completion(client, job["id"])
    assert job["status"] == "PARTIAL" and job["error_count"] == 1
    assert job["processed_rows"] == len(expects)
    assert job["output_available"] and job["output_filename"] == "my_streamers_processed.xlsx"

    rows = client.get(f"/api/jobs/{job['id']}/rows?limit=1000").json()
    assert [r["original_row"] for r in rows["items"]] == list(range(2, len(expects) + 2))
    detail = client.get(f"/api/jobs/{job['id']}/rows/2").json()
    assert detail["decision"] == "MATCH" and detail["evidence_json"]["candidates"]

    d = client.get(f"/api/jobs/{job['id']}/download")
    assert d.status_code == 200 and "my_streamers_processed.xlsx" in d.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(d.content))["Streamers"]
    assert ws["A2"].value == "starwraith" and ws["C2"].value == "StarWraith"
    rr = client.get(f"/api/jobs/{job['id']}/review-report")
    assert rr.status_code == 200 and "my_streamers_review.xlsx" in rr.headers["content-disposition"]
    assert set(load_workbook(io.BytesIO(rr.content)).sheetnames) == {"Summary", "Rows", "Candidates"}

    # --- review workflow: confirm the ambiguous "twinz" row, reject "conflicted"
    items = client.get(f"/api/jobs/{job['id']}/reviews").json()
    by_src = {i["source_value"]: i for i in items}
    assert {"twinz", "conflicted"} <= set(by_src)
    twinz_row = by_src["twinz"]["original_row"]
    assert by_src["twinz"]["source_profile"]["username"] == "twinz" and by_src["twinz"]["candidate"]

    bad = client.post(
        f"/api/jobs/{job['id']}/reviews",
        json={"original_row": twinz_row, "verdict": "CONFIRM", "target_username": "invented_name"},
    )
    assert bad.status_code == 400  # never accept an id the engine did not discover

    ok = client.post(
        f"/api/jobs/{job['id']}/reviews",
        json={"original_row": twinz_row, "verdict": "CONFIRM", "target_username": "twinz"},
    )
    assert ok.status_code == 200, ok.text
    conf_row = by_src["conflicted"]["original_row"]
    rej = client.post(
        f"/api/jobs/{job['id']}/reviews",
        json={"original_row": conf_row, "verdict": "REJECT", "target_username": "conflicted"},
    )
    assert rej.status_code == 200

    job2 = client.get(f"/api/jobs/{job['id']}").json()
    assert job2["match_count"] == job["match_count"] + 1 and job2["export_stale"]
    ws = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download").content))["Streamers"]
    assert ws.cell(row=twinz_row, column=3).value == "twinz"
    assert ws.cell(row=conf_row, column=3).value is None
    assert client.get(f"/api/jobs/{job['id']}").json()["verification_json"]["ok"]


def test_ambiguous_workbook_requires_platform_choice(client, fake, tmp_path):
    fake.add_twitch("abcdef")
    path = make_workbook(
        tmp_path / "amb.xlsx",
        KICK_HEADERS,
        [[None, "France", "abcdef", None], [None, "Spain", "ghijkl", None]],
    )
    job = upload(client, path).json()
    assert job["needs_platform_choice"] and job["detected_platform"] is None
    r = client.post(f"/api/jobs/{job['id']}/start", json={})
    assert r.status_code == 409
    done = run_to_completion(client, job["id"], "twitch")
    assert done["source_platform"] == "twitch" and done["status"] == "COMPLETED"


def test_twitch_workbook_via_api(client, fake, tmp_path):
    fake.add_twitch("kosstochka")
    path = make_workbook(
        tmp_path / "t.xlsx",
        TWITCH_HEADERS,
        [["kosstochka", "Italy", "None", None], ["unknowntw", "Italy", None, None]],
    )
    job = run_to_completion(client, upload(client, path).json()["id"])
    ws = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download").content))["Streamers"]
    assert (ws["C2"].value, ws["D2"].value) == (None, "no kick id")  # "None" placeholder cleared
    assert (ws["C3"].value, ws["D3"].value) == (None, "no Id on both platforms")


@pytest.mark.parametrize(
    ("name", "content", "code"),
    [
        ("x.csv", b"a,b", 400),
        ("x.xlsx", b"", 400),
        ("x.xlsx", b"not a workbook", 422),
    ],
)
def test_upload_validation(client, name, content, code):
    r = client.post("/api/jobs", files={"file": (name, content, XLSX)})
    assert r.status_code == code


def test_download_before_completion_rejected(client, fake, tmp_path):
    path = make_workbook(tmp_path / "w.xlsx", KICK_HEADERS, [["abc", "France", None, None]])
    job = upload(client, path).json()
    assert client.get(f"/api/jobs/{job['id']}/download").status_code == 409


def test_retry_failed_endpoint(client, fake, tmp_path):
    fake.add_kick("brokenrow", image=None)
    fake.fail(lambda r: r.url.path == "/public/v1/channels", status=503)
    path = make_workbook(tmp_path / "w.xlsx", KICK_HEADERS, [["brokenrow", "France", None, None]])
    job = run_to_completion(client, upload(client, path).json()["id"])
    assert job["status"] == "PARTIAL" and job["error_count"] == 1
    ws = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download").content))["Streamers"]
    assert ws["C2"].value is None and ws["D2"].value is None  # API failure never becomes a remark
    fake.failures.clear()
    r = client.post(f"/api/jobs/{job['id']}/retry-failed")
    assert r.status_code == 200
    wait_inline(job["id"], timeout=60)
    job = client.get(f"/api/jobs/{job['id']}").json()
    assert job["status"] == "COMPLETED" and job["error_count"] == 0


def test_outdated_export_is_rebuilt_on_download(client, fake, sf, tmp_path):
    from app.models import ProcessingJob

    fake.add_twitch("kosstochka")
    path = make_workbook(tmp_path / "old.xlsx", TWITCH_HEADERS, [["kosstochka", "Italy", None, None]])
    job = run_to_completion(client, upload(client, path).json()["id"])
    with sf() as s:  # simulate a file produced by an older export layout
        j = s.get(ProcessingJob, job["id"])
        j.verification_json = {**j.verification_json, "export_format": 1}
        s.commit()
    ws = load_workbook(io.BytesIO(client.get(f"/api/jobs/{job['id']}/download").content))["Streamers"]
    assert ws["E2"].value == "https://www.twitch.tv/kosstochka"
    assert client.get(f"/api/jobs/{job['id']}").json()["verification_json"]["export_format"] == 3
