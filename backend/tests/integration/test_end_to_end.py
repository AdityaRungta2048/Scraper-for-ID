"""End-to-end: workbook -> detection -> real adapters (fake HTTP) -> matching -> export -> verification."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.models import JobStatus, MatchDecision, ProcessingJob, ProcessingRow
from app.services.jobs import create_job, start_job
from app.version import MATCHING_ENGINE_VERSION
from app.workers.processor import process_job
from tests.excel_helpers import KICK_HEADERS, TWITCH_HEADERS, make_workbook
from tests.scenarios import Expect, build_kick_world, build_twitch_world


def _book(tmp_path: Path, name: str, headers, expects: list[Expect]) -> Path:
    return make_workbook(tmp_path / name, headers, [[e.source, e.country, None, None] for e in expects])


async def run(sf, path: Path, platform: str | None = None) -> ProcessingJob:
    settings = get_settings()
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, platform)
        job_id = job.id
    await process_job(job_id, settings=settings, sf=sf)
    with sf() as s:
        return s.get(ProcessingJob, job_id)


def rows_by_number(sf, job_id):
    with sf() as s:
        return {
            r.original_row: r
            for r in s.execute(select(ProcessingRow).where(ProcessingRow.job_id == job_id)).scalars()
        }


def assert_output(path: Path, expects: list[Expect], dest_col: int, src_col: int):
    ws = load_workbook(path)["Streamers"]
    for i, e in enumerate(expects, start=2):
        assert ws.cell(row=i, column=src_col).value == e.source, f"row {i} source changed"
        assert ws.cell(row=i, column=2).value == e.country, f"row {i} country changed"
        assert ws.cell(row=i, column=dest_col).value == e.dest, f"row {i} ({e.case}) destination"
        assert ws.cell(row=i, column=4).value == e.remarks, f"row {i} ({e.case}) remarks"
    assert ws.max_row == len(expects) + 1


async def test_kick_source_workbook_end_to_end(fake, sf, tmp_path):
    expects = build_kick_world(fake)
    job = await run(sf, _book(tmp_path, "streamers.xlsx", KICK_HEADERS, expects))
    assert job.source_platform == "kick"
    rows = rows_by_number(sf, job.id)
    problems = []
    for i, e in enumerate(expects, start=2):
        r = rows[i]
        if (
            (e.status != "*" and r.status != e.status)
            or r.output_destination != e.dest
            or ((r.output_remarks if r.write_remarks else None) != e.remarks)
        ):
            problems.append(
                f"row {i} [{e.case}] {e.source}: status={r.status} dest={r.output_destination} "
                f"remarks={r.output_remarks} reason={r.reason}"
            )
        if e.status == "*":
            assert r.status in ("NO_MATCH", "REVIEW"), (e.case, r.status, r.reason)
    assert not problems, "\n".join(problems)
    # one API failure row -> PARTIAL, but a verified workbook is still produced
    assert job.status == JobStatus.PARTIAL.value and job.error_count == 1
    assert job.verification_json["ok"], job.verification_json
    assert Path(job.output_path).name == "streamers_processed.xlsx"
    assert_output(Path(job.output_path), expects, dest_col=3, src_col=1)
    # every automatic decision is explainable and versioned
    with sf() as s:
        decisions = list(s.execute(select(MatchDecision)).scalars())
    assert decisions and all(
        d.matching_engine_version == MATCHING_ENGINE_VERSION and d.reason for d in decisions
    )
    match_row = rows[2]
    best = next(c for c in match_row.evidence_json["candidates"] if c["username"] == "starwraith")
    assert best["evidence"]["profile_image_similarity"] >= 0.9 and best["reason"]


async def test_twitch_source_workbook_end_to_end(fake, sf, tmp_path, monkeypatch):
    monkeypatch.setenv("BRAVE_SEARCH_API_KEY", "test-key")
    get_settings.cache_clear()
    expects = build_twitch_world(fake)
    job = await run(sf, _book(tmp_path, "twitch_list.xlsx", TWITCH_HEADERS, expects))
    assert job.source_platform == "twitch"
    rows = rows_by_number(sf, job.id)
    for i, e in enumerate(expects, start=2):
        r = rows[i]
        assert r.output_destination == e.dest, (e.case, r.status, r.reason)
        assert (r.output_remarks if r.write_remarks else None) == e.remarks, (e.case, r.status, r.reason)
        if e.status != "*":
            assert r.status == e.status, (e.case, r.status, r.reason)
    assert job.status == JobStatus.COMPLETED.value
    assert_output(Path(job.output_path), expects, dest_col=3, src_col=1)
    assert Path(job.review_path).name == "twitch_list_review.xlsx"


async def test_duplicates_resolved_once(fake, sf, tmp_path):
    build_kick_world(fake)
    expects = [Expect("starwraith", "Spain", "StarWraith", None, "MATCH", "dup")] * 6
    job = await run(sf, _book(tmp_path, "dups.xlsx", KICK_HEADERS, expects))
    assert job.match_count == 6 and job.total_rows == 6
    # existence of the source checked once; image of each account fetched once
    assert fake.calls["api.kick.com/public/v1/channels"] <= 2
    assert sum(v for k, v in fake.calls.items() if "jtv_user_pictures" in k) == 1


async def test_retry_failed_rows_after_outage(fake, sf, tmp_path):
    expects = build_kick_world(fake)
    job = await run(sf, _book(tmp_path, "s.xlsx", KICK_HEADERS, expects))
    assert job.status == "PARTIAL"
    fake.failures.clear()
    fake.add_kick("flaky", description="instagram.com/flaky.creator", image=to_img(90))
    fake.add_twitch("flaky", description="IG: @flaky.creator", image=to_img_v(90))
    from app.models import RowStatus
    from app.models.tables import ERROR_STATUSES

    with sf() as s:
        for r in s.execute(select(ProcessingRow).where(ProcessingRow.job_id == job.id)).scalars():
            if RowStatus(r.status) in ERROR_STATUSES:
                r.status = RowStatus.PENDING.value
        j = s.get(ProcessingJob, job.id)
        j.status = JobStatus.QUEUED.value
        s.commit()
    before = fake.api_calls("api.twitch.tv")
    await process_job(job.id, settings=get_settings(), sf=sf)
    with sf() as s:
        j = s.get(ProcessingJob, job.id)
        assert j.status == "COMPLETED" and j.error_count == 0
        flaky_row = s.execute(
            select(ProcessingRow).where(ProcessingRow.job_id == job.id, ProcessingRow.source_value == "flaky")
        ).scalar_one()
        assert flaky_row.output_destination == "flaky"
    # only the failed row was reprocessed (others were not repeated)
    assert fake.api_calls("api.twitch.tv") - before <= 6
    ws = load_workbook(j.output_path)["Streamers"]
    assert ws.cell(row=14, column=1).value == "flaky" and ws.cell(row=14, column=3).value == "flaky"


def to_img(seed):
    from tests.scenarios import img

    return img(seed)


def to_img_v(seed):
    from tests.scenarios import img_v

    return img_v(seed)


async def test_missing_credentials_fail_job_clearly(fake, sf, tmp_path, monkeypatch):
    monkeypatch.setenv("KICK_CLIENT_SECRET", "")
    get_settings.cache_clear()
    job = await run(
        sf, _book(tmp_path, "c.xlsx", KICK_HEADERS, [Expect("abc", "France", None, None, "*", "")])
    )
    assert job.status == "FAILED" and "credentials" in job.error_message.lower()
    assert job.output_path is None


async def test_auth_failure_is_config_error_not_no_match(fake, sf, tmp_path):
    fake.fail(lambda r: r.url.path == "/oauth2/token", status=401)
    job = await run(
        sf, _book(tmp_path, "c.xlsx", TWITCH_HEADERS, [Expect("abc", "France", None, None, "*", "")])
    )
    assert job.status == "FAILED"
    rows = rows_by_number(sf, job.id)
    assert rows[2].status == "API_ERROR" and not rows[2].write_remarks


@pytest.mark.parametrize("policy", ["preserve", "overwrite"])
async def test_existing_destination_policy(fake, sf, tmp_path, monkeypatch, policy):
    monkeypatch.setenv("EXISTING_DESTINATION_POLICY", policy)
    get_settings.cache_clear()
    build_kick_world(fake)
    path = make_workbook(
        tmp_path / "pre.xlsx", KICK_HEADERS, [["starwraith", "Spain", "typed_by_hand", "keep me"]]
    )
    job = await run(sf, path)
    ws = load_workbook(job.output_path)["Streamers"]
    if policy == "preserve":
        assert ws["C2"].value == "typed_by_hand" and ws["D2"].value == "keep me"
    else:
        assert ws["C2"].value == "StarWraith" and ws["D2"].value == "keep me"


async def test_channel_link_columns(fake, sf, tmp_path):
    expects = build_kick_world(fake)
    job = await run(sf, _book(tmp_path, "links.xlsx", KICK_HEADERS, expects))
    ws = load_workbook(job.output_path)["Streamers"]
    assert [ws.cell(row=1, column=c).value for c in (5, 6)] == ["twitch_id_link", "kick_id_link"]
    row = {e.source: i for i, e in enumerate(expects, start=2) if e.source}

    def cell(src, col):
        return ws.cell(row=row[src], column=col)

    # confident match: plain, clickable channel links on both platforms
    assert cell("starwraith", 5).value == "https://www.twitch.tv/starwraith"
    assert cell("starwraith", 5).hyperlink.target == "https://www.twitch.tv/starwraith"
    assert cell("starwraith", 6).value == "https://kick.com/starwraith"
    # exactly one link per cell, never a list or labels
    for i in range(2, len(expects) + 2):
        for col in (5, 6):
            v = ws.cell(row=i, column=col).value
            assert v is None or ("\n" not in v and " " not in v), (i, v)
    # review row: the review candidate; the ID column stays empty
    assert cell("twinz", 5).value == "https://www.twitch.tv/twinz" and cell("twinz", 3).value is None
    # no match: the closest candidate (same name first) is linked, the ID stays empty
    assert cell("alex123", 5).value == "https://www.twitch.tv/alex123" and cell("alex123", 3).value is None
    # source missing: the same-name account is linked; no source link
    assert cell("ghostkick", 5).value == "https://www.twitch.tv/ghostkick"
    assert cell("ghostkick", 6).value is None
    # nothing found / API error rows get no links
    assert cell("unknownabc", 5).value is None and cell("flaky", 5).value is None
    assert job.verification_json["ok"]


async def test_existing_link_columns_are_reused(fake, sf, tmp_path):
    build_kick_world(fake)
    headers = ["id_kick", "country", "id_twitch", "remarks", "Twitch ID Link", "Kick ID Link"]
    path = make_workbook(
        tmp_path / "again.xlsx", headers, [["starwraith", "Spain", None, None, "stale", None]]
    )
    job = await run(sf, path)
    ws = load_workbook(job.output_path)["Streamers"]
    assert ws.max_column == 6  # no extra columns appended
    assert ws["E2"].value == "https://www.twitch.tv/starwraith"  # stale value replaced


async def test_jobs_from_before_link_columns_get_them_on_next_export(fake, sf, tmp_path):
    from app.services.jobs import build_export

    build_kick_world(fake)
    job = await run(
        sf, _book(tmp_path, "old.xlsx", KICK_HEADERS, [Expect("starwraith", "Spain", None, None, "*", "")])
    )
    with sf() as s:
        j = s.get(ProcessingJob, job.id)
        j.columns_json = {
            k: v for k, v in j.columns_json.items() if k not in ("twitch_link", "kick_link", "new_headers")
        }
        s.commit()
        report = build_export(s, get_settings(), j)
        assert report["ok"]
        ws = load_workbook(j.output_path)["Streamers"]
    assert ws["E1"].value == "twitch_id_link" and ws["E2"].value == "https://www.twitch.tv/starwraith"
