"""Output-integrity tests (spec §74): asynchronous completion order never changes the workbook."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.models import JobStatus, ProcessingJob, ProcessingRow
from app.services.context import MatchingContext
from app.services.jobs import create_job, start_job
from app.version import MATCHING_ENGINE_VERSION
from app.workers.processor import process_job
from tests.excel_helpers import KICK_HEADERS, make_workbook


class _NoPlatforms:
    async def aclose(self) -> None:
        return None


class OrchestratedResolver:
    """Finishes identities in a forced order, regardless of the order they were started."""

    def __init__(self, completion_order: list[str]) -> None:
        self.order = completion_order
        self.done = {v: asyncio.Event() for v in completion_order}
        self.started: list[str] = []

    async def resolve(self, platform: str, value: str, country: str | None, rejected=None) -> dict[str, Any]:
        self.started.append(value)
        idx = self.order.index(value)
        if idx > 0:
            await asyncio.wait_for(self.done[self.order[idx - 1]].wait(), timeout=10)
        return {
            "engine_version": MATCHING_ENGINE_VERSION,
            "source_platform": platform,
            "target_platform": "twitch",
            "source_value": value,
            "source_key": value.lower(),
            "source_status": "EXISTS",
            "decision": "MATCH",
            "target_status": None,
            "matched_id": f"{value}_TW",
            "matched_username": f"{value.lower()}_tw",
            "review_candidate": None,
            "confidence": 99.0,
            "reason": "test",
            "candidates": [],
            "from_cache": True,
        }


def _run(sf, tmp_path: Path, values: list[str | None], completion: list[str], name="order.xlsx"):
    settings = get_settings()
    path = make_workbook(tmp_path / name, KICK_HEADERS, [[v, "France", None, None] for v in values])
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, "kick")
        job_id = job.id
    resolver = OrchestratedResolver(completion)
    finished: list[str] = []

    def ctx_factory(_settings, _sf):
        return MatchingContext(resolver=resolver, platforms=_NoPlatforms(), cache=None)  # type: ignore[arg-type]

    def on_done(value: str) -> None:
        finished.append(value)
        resolver.done[value].set()

    asyncio.run(
        process_job(job_id, settings=settings, sf=sf, context_factory=ctx_factory, on_group_done=on_done)
    )
    with sf() as s:
        job = s.get(ProcessingJob, job_id)
    return job, resolver, finished


def test_async_completion_order_does_not_change_output(env, sf, tmp_path):
    job, resolver, finished = _run(sf, tmp_path, ["A", "B", "C", "D"], completion=["B", "D", "A", "C"])
    assert resolver.started[:4] == ["A", "B", "C", "D"]  # all started concurrently in row order
    assert finished == ["B", "D", "A", "C"]  # ...but finished B, D, A, C
    assert job.status == JobStatus.COMPLETED.value and job.verification_json["ok"]
    ws = load_workbook(job.output_path)["Streamers"]
    assert [ws.cell(row=r, column=1).value for r in range(2, 6)] == ["A", "B", "C", "D"]
    assert [ws.cell(row=r, column=3).value for r in range(2, 6)] == ["A_TW", "B_TW", "C_TW", "D_TW"]


def test_many_rows_with_duplicates_keep_positions(env, sf, tmp_path):
    values = [f"u{i % 37}" for i in range(300)]
    order = sorted(set(values), key=lambda v: -int(v[1:]))  # reverse completion
    get_settings().row_concurrency = 64  # all identities in flight so the forced order is reachable
    job, _, finished = _run(sf, tmp_path, values, completion=order, name="big.xlsx")
    assert len(finished) == 37  # duplicates resolved once
    ws = load_workbook(job.output_path)["Streamers"]
    for i, v in enumerate(values, start=2):
        assert ws.cell(row=i, column=1).value == v
        assert ws.cell(row=i, column=3).value == f"{v}_TW"
    assert ws.max_row == 301


def test_resume_processes_only_remaining_rows(env, sf, tmp_path):
    settings = get_settings()
    path = make_workbook(tmp_path / "r.xlsx", KICK_HEADERS, [[v, "France", None, None] for v in "ABCD"])
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, "kick")
        job_id = job.id
        # simulate a crash after rows 2 and 3 were persisted
        for r in s.execute(select(ProcessingRow).where(ProcessingRow.job_id == job_id)).scalars():
            if r.original_row in (2, 3):
                r.status, r.decision, r.matched_id = "MATCH", "MATCH", f"{r.source_value}_TW"
                r.output_destination, r.write_destination = f"{r.source_value}_TW", True
                r.evidence_json = {"source_status": "EXISTS"}
        s.commit()
    resolver = OrchestratedResolver(["C", "D"])

    def ctx_factory(_settings, _sf):
        return MatchingContext(resolver=resolver, platforms=_NoPlatforms(), cache=None)  # type: ignore[arg-type]

    asyncio.run(
        process_job(
            job_id,
            settings=settings,
            sf=sf,
            context_factory=ctx_factory,
            on_group_done=lambda v: resolver.done[v].set(),
        )
    )
    assert resolver.started == ["C", "D"]  # A and B were not repeated
    with sf() as s:
        job = s.get(ProcessingJob, job_id)
    ws = load_workbook(job.output_path)["Streamers"]
    assert [ws.cell(row=r, column=3).value for r in range(2, 6)] == ["A_TW", "B_TW", "C_TW", "D_TW"]


def test_job_cannot_be_claimed_twice(env, sf, tmp_path):
    from app.workers.processor import claim_job

    settings = get_settings()
    path = make_workbook(tmp_path / "c.xlsx", KICK_HEADERS, [["A", "France", None, None]])
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, "kick")
    assert claim_job(sf, job.id) is True
    assert claim_job(sf, job.id) is False


def test_cancelled_job_stops(env, sf, tmp_path):
    settings = get_settings()
    path = make_workbook(tmp_path / "x.xlsx", KICK_HEADERS, [[v, "France", None, None] for v in "ABC"])
    with sf() as s:
        job, _ = create_job(s, settings, path.name, path.read_bytes())
        start_job(s, job, "kick")
        job_id = job.id

    class CancellingResolver(OrchestratedResolver):
        async def resolve(self, platform, value, country, rejected=None):
            with sf() as s2:
                j = s2.get(ProcessingJob, job_id)
                j.status = JobStatus.CANCELLED.value
                s2.commit()
            return await super().resolve(platform, value, country)

    resolver = CancellingResolver(["A", "B", "C"])
    settings.row_concurrency = 1
    asyncio.run(
        process_job(
            job_id,
            settings=settings,
            sf=sf,
            context_factory=lambda *_: MatchingContext(
                resolver=resolver, platforms=_NoPlatforms(), cache=None
            ),  # type: ignore[arg-type]
            on_group_done=lambda v: resolver.done[v].set(),
        )
    )
    with sf() as s:
        job = s.get(ProcessingJob, job_id)
        assert job.status == "CANCELLED" and job.output_path is None
    assert resolver.started == ["A"]
