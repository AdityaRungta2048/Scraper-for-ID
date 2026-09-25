"""Secondary deliverable: <name>_review.xlsx with per-row decisions and candidate evidence."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DECISION_FILLS = {
    "MATCH": PatternFill("solid", fgColor="D1FAE5"),
    "REVIEW": PatternFill("solid", fgColor="FEF3C7"),
    "NO_MATCH": PatternFill("solid", fgColor="F3F4F6"),
}


def review_filename(original_name: str) -> str:
    return f"{Path(original_name).stem}_review.xlsx"


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_review_report(path: Path, job: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    for k, v in [
        ("File", job.get("filename")),
        ("Source platform", job.get("source_platform")),
        ("Rows", job.get("total_rows")),
        ("Matches", job.get("match_count")),
        ("No match / not found", job.get("no_match_count")),
        ("Review", job.get("review_count")),
        ("Errors", job.get("error_count")),
        ("Matching engine version", job.get("matching_engine_version")),
        (
            "Note",
            "The primary deliverable is the *_processed.xlsx workbook. This report is supporting evidence.",
        ),
    ]:
        summary.append([k, v])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 80

    target = "twitch" if job.get("source_platform") == "kick" else "kick"
    main_rows, cand_rows = [], []
    for r in rows:
        ev = r.get("evidence_json") or {}
        main_rows.append(
            [
                r["original_row"],
                job.get("source_platform"),
                r.get("source_value"),
                r.get("country"),
                r.get("status"),
                r.get("decision"),
                r.get("confidence"),
                target,
                r.get("matched_id") or r.get("review_candidate"),
                r.get("output_destination"),
                r.get("output_remarks"),
                r.get("manual_verdict"),
                r.get("reason") or r.get("error_message"),
                json.dumps((ev.get("best") or {}).get("evidence") or {}, ensure_ascii=False),
            ]
        )
        for c in ev.get("candidates", []) or []:
            e = c.get("evidence", {})
            cand_rows.append(
                [
                    r["original_row"],
                    r.get("source_value"),
                    c.get("platform"),
                    c.get("username"),
                    c.get("decision"),
                    c.get("confidence"),
                    e.get("username_similarity"),
                    e.get("display_name_similarity"),
                    e.get("profile_image_similarity"),
                    e.get("explicit_cross_link"),
                    ", ".join(e.get("shared_social_accounts") or []),
                    e.get("bio_similarity"),
                    e.get("content_similarity"),
                    e.get("country_match"),
                    ", ".join(c.get("discovered_via") or []),
                    c.get("reason"),
                ]
            )
    _sheet(
        wb,
        "Rows",
        [
            "original_row",
            "source_platform",
            "source_id",
            "country",
            "status",
            "decision",
            "confidence",
            "candidate_platform",
            "candidate_id",
            "written_destination",
            "written_remarks",
            "manual_verdict",
            "reason",
            "evidence",
        ],
        main_rows,
        [12, 14, 22, 12, 18, 11, 11, 16, 22, 22, 24, 14, 90, 60],
    )
    _sheet(
        wb,
        "Candidates",
        [
            "original_row",
            "source_id",
            "candidate_platform",
            "candidate_id",
            "decision",
            "confidence",
            "username_sim",
            "display_name_sim",
            "image_sim",
            "explicit_link",
            "shared_socials",
            "bio_sim",
            "content_sim",
            "country_match",
            "discovered_via",
            "reason",
        ],
        cand_rows,
        [12, 22, 16, 22, 11, 11, 12, 14, 10, 12, 30, 9, 11, 12, 28, 90],
    )
    ws = wb["Rows"]
    for row in ws.iter_rows(min_row=2):
        fill = DECISION_FILLS.get(str(row[5].value))
        if fill:
            for cell in row[:7]:
                cell.fill = fill
        row[12].alignment = Alignment(wrap_text=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
