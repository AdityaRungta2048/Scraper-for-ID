"""ExcelImporter: validation, header location, source-platform detection, row extraction.

Each data row keeps its immutable 1-based Excel row number (``original_row``).
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.excel.state_machine import is_blank

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
HEADER_SCAN_ROWS = 20
MAX_DATA_ROWS = 200_000

_ALIASES = {
    "id_kick": "id_kick",
    "kick_id": "id_kick",
    "idkick": "id_kick",
    "kick": "id_kick",
    "id_twitch": "id_twitch",
    "twitch_id": "id_twitch",
    "idtwitch": "id_twitch",
    "twitch": "id_twitch",
    "country": "country",
    "pays": "country",
    "land": "country",
    "pais": "country",
    "paese": "country",
    "remarks": "remarks",
    "remark": "remarks",
    "comments": "remarks",
    "comment": "remarks",
    "notes": "remarks",
    "twitch_id_link": "twitch_id_link",
    "twitch_link": "twitch_id_link",
    "twitch_url": "twitch_id_link",
    "twitch_channel_link": "twitch_id_link",
    "kick_id_link": "kick_id_link",
    "kick_link": "kick_id_link",
    "kick_url": "kick_id_link",
    "kick_channel_link": "kick_id_link",
}
LINK_HEADERS = {"twitch": "twitch_id_link", "kick": "kick_id_link"}


class WorkbookValidationError(ValueError):
    pass


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[\s\-\.]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return _ALIASES.get(text, text)


@dataclass
class ColumnMap:
    sheet_name: str
    header_row: int
    id_kick: int
    id_twitch: int
    country: int | None
    remarks: int | None  # None => the exporter adds a "remarks" header in ``remarks_new_col``
    remarks_new_col: int | None = None
    # Channel-link columns (existing, or appended after the last used column).
    twitch_link: int | None = None
    kick_link: int | None = None
    # header cells the exporter must create: {"<column index>": "<header text>"}
    new_headers: dict[str, str] | None = None

    def source_col(self, platform: str) -> int:
        return self.id_kick if platform == "kick" else self.id_twitch

    def dest_col(self, platform: str) -> int:
        return self.id_twitch if platform == "kick" else self.id_kick

    @property
    def remarks_col(self) -> int:
        col = self.remarks if self.remarks is not None else self.remarks_new_col
        assert col is not None
        return col

    def link_col(self, platform: str) -> int | None:
        return self.twitch_link if platform == "twitch" else self.kick_link

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict[str, Any]) -> ColumnMap:
        return cls(**{k: d.get(k) for k in cls.__dataclass_fields__})  # type: ignore[arg-type]


@dataclass
class WorkbookAnalysis:
    sheet_name: str
    header_row: int
    headers: list[str]
    columns: ColumnMap
    detected_platform: str | None
    ambiguous: bool
    detection_reason: str
    row_count: int
    data_rows: int
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d["columns"] = self.columns.to_dict()
        return d


@dataclass
class ImportedRow:
    original_row: int
    source_value: str | None
    country: str | None
    existing_destination: str | None
    existing_remarks: str | None


def cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    return text if text.strip() else None


def _open(path: Path, **kwargs: Any) -> Any:
    if path.suffix.lower() == ".xls":
        raise WorkbookValidationError("Legacy .xls files are not supported. Please save the file as .xlsx.")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise WorkbookValidationError(
            f"Unsupported file type '{path.suffix}'. Upload an .xlsx or .xlsm workbook."
        )
    if not zipfile.is_zipfile(path):
        raise WorkbookValidationError("The file is not a valid Excel workbook (corrupted or wrong format).")
    try:
        return load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm", **kwargs)
    except Exception as exc:
        raise WorkbookValidationError(
            f"The workbook could not be opened: {type(exc).__name__}: {exc}"
        ) from exc


def analyze_workbook(path: Path) -> WorkbookAnalysis:
    wb = _open(path, data_only=True)
    if not wb.worksheets:
        raise WorkbookValidationError("The workbook contains no worksheets.")
    matches: list[tuple[Any, int, dict[str, int], list[str]]] = []
    partial: list[str] = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible" and len(wb.worksheets) > 1:
            continue
        for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
            raw = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            normed = [normalize_header(v) for v in raw]
            cols: dict[str, int] = {}
            for idx, h in enumerate(normed, start=1):
                if (
                    h in {"id_kick", "id_twitch", "country", "remarks", *LINK_HEADERS.values()}
                    and h not in cols
                ):
                    cols[h] = idx
            if "id_kick" in cols and "id_twitch" in cols:
                matches.append((ws, r, cols, [str(v) if v is not None else "" for v in raw]))
                break
            if "id_kick" in cols or "id_twitch" in cols:
                partial.append(
                    f"sheet '{ws.title}' row {r} has only {'id_kick' if 'id_kick' in cols else 'id_twitch'}"
                )
    if not matches:
        detail = ("; ".join(partial) + ". ") if partial else ""
        raise WorkbookValidationError(
            f"Required headers not found. {detail}Expected 'id_kick | country | id_twitch | remarks' "
            "(Kick source) or 'id_twitch | country | id_kick | remarks' (Twitch source) in the first "
            f"{HEADER_SCAN_ROWS} rows."
        )
    warnings: list[str] = []
    if len(matches) > 1:
        warnings.append(
            f"Several sheets contain the expected headers; using '{matches[0][0].title}'. "
            "Other sheets are preserved unchanged."
        )
    ws, header_row, cols, headers = matches[0]
    remarks_new_col = None
    if "remarks" not in cols:
        remarks_new_col = ws.max_column + 1
        warnings.append(
            f"No 'remarks' column found; a 'remarks' header will be added in column "
            f"{get_column_letter(remarks_new_col)}."
        )
    if "country" not in cols:
        warnings.append("No 'country' column found; country evidence will not be used.")
    link_cols, new_headers = _link_columns(
        ws, header_row, cols, reserved={remarks_new_col} if remarks_new_col else set()
    )
    colmap = ColumnMap(
        sheet_name=ws.title,
        header_row=header_row,
        id_kick=cols["id_kick"],
        id_twitch=cols["id_twitch"],
        country=cols.get("country"),
        remarks=cols.get("remarks"),
        remarks_new_col=remarks_new_col,
        twitch_link=link_cols["twitch"],
        kick_link=link_cols["kick"],
        new_headers=new_headers,
    )

    last = _last_data_row(ws, header_row, [c for c in cols.values()])
    data_rows = max(0, last - header_row)
    if data_rows > MAX_DATA_ROWS:
        raise WorkbookValidationError(f"Too many rows ({data_rows}); the limit is {MAX_DATA_ROWS}.")
    kick_filled = twitch_filled = 0
    for r in range(header_row + 1, last + 1):
        if not is_blank(ws.cell(row=r, column=colmap.id_kick).value):
            kick_filled += 1
        if not is_blank(ws.cell(row=r, column=colmap.id_twitch).value):
            twitch_filled += 1
    if data_rows == 0 or (kick_filled == 0 and twitch_filled == 0):
        raise WorkbookValidationError("The workbook has the expected headers but no data rows.")

    by_order = "kick" if colmap.id_kick < colmap.id_twitch else "twitch"
    fk, ft = kick_filled / data_rows, twitch_filled / data_rows
    by_fill = "kick" if fk > ft + 0.2 else ("twitch" if ft > fk + 0.2 else None)
    if by_fill is None or by_fill == by_order:
        detected, ambiguous = by_order, False
        reason = (
            f"'{'id_kick' if by_order == 'kick' else 'id_twitch'}' is the first ID column "
            f"(fill: id_kick {fk:.0%}, id_twitch {ft:.0%})."
        )
    else:
        detected, ambiguous = None, True
        reason = (
            f"Column order suggests a {by_order.upper()} source, but fill rates suggest {by_fill.upper()} "
            f"(id_kick {fk:.0%} filled, id_twitch {ft:.0%} filled). Please choose the source platform."
        )

    for r in range(header_row + 1, last + 1):
        src = ws.cell(row=r, column=colmap.source_col(detected or by_order)).value
        if isinstance(src, str) and (" " in src.strip() or len(src.strip()) > 64):
            warnings.append(f"Row {r}: source value {src!r} looks malformed (spaces/too long).")
            if len(warnings) > 25:
                warnings.append("…more warnings suppressed.")
                break
    return WorkbookAnalysis(
        sheet_name=ws.title,
        header_row=header_row,
        headers=headers,
        columns=colmap,
        detected_platform=detected,
        ambiguous=ambiguous,
        detection_reason=reason,
        row_count=data_rows,
        data_rows=data_rows,
        warnings=warnings,
    )


def _last_data_row(ws: Any, header_row: int, cols: list[int]) -> int:
    last = header_row
    for r in range(ws.max_row, header_row, -1):
        if any(not is_blank(ws.cell(row=r, column=c).value) for c in cols):
            last = r
            break
    return last


def read_rows(path: Path, columns: ColumnMap, source_platform: str) -> list[ImportedRow]:
    wb = _open(path, data_only=True)
    ws = wb[columns.sheet_name]
    cols = [columns.id_kick, columns.id_twitch] + [c for c in (columns.country, columns.remarks) if c]
    last = _last_data_row(ws, columns.header_row, cols)
    src_col, dst_col = columns.source_col(source_platform), columns.dest_col(source_platform)
    rows: list[ImportedRow] = []
    for r in range(columns.header_row + 1, last + 1):
        rows.append(
            ImportedRow(
                original_row=r,
                source_value=cell_text(ws.cell(row=r, column=src_col).value),
                country=cell_text(ws.cell(row=r, column=columns.country).value) if columns.country else None,
                existing_destination=cell_text(ws.cell(row=r, column=dst_col).value),
                existing_remarks=cell_text(ws.cell(row=r, column=columns.remarks).value)
                if columns.remarks
                else None,
            )
        )
    return rows


def _link_columns(
    ws: Any, header_row: int, cols: dict[str, int], reserved: set[int]
) -> tuple[dict[str, int], dict[str, str]]:
    """Reuse existing twitch_id_link / kick_id_link columns, otherwise place new ones in the
    first completely empty columns to the right of all existing data (never between columns)."""
    result: dict[str, int] = {}
    new_headers: dict[str, str] = {}
    last_used = max(
        [c for c in range(1, ws.max_column + 1) if _column_has_data(ws, c)] + list(reserved) + [0]
    )
    next_col = last_used + 1
    for platform in ("twitch", "kick"):
        header = LINK_HEADERS[platform]
        if header in cols:
            result[platform] = cols[header]
            continue
        result[platform] = next_col
        new_headers[str(next_col)] = header
        next_col += 1
    return result, new_headers


def _column_has_data(ws: Any, col: int) -> bool:
    return any(ws.cell(row=r, column=col).value not in (None, "") for r in range(1, ws.max_row + 1))


def upgrade_columns(columns: ColumnMap, path: Path) -> ColumnMap:
    """Jobs created before link columns existed: compute their placement from the original file."""
    if columns.twitch_link is not None and columns.kick_link is not None:
        return columns
    fresh = analyze_workbook(path).columns
    columns.twitch_link, columns.kick_link, columns.new_headers = (
        fresh.twitch_link,
        fresh.kick_link,
        fresh.new_headers,
    )
    return columns
