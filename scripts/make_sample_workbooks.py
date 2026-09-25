"""Create sample input workbooks with the same structure as the project brief.

    python scripts/make_sample_workbooks.py        # writes samples/*.xlsx

The destination/remarks columns are intentionally EMPTY: they are what the app fills in.
The IDs are the examples from the brief; nothing about them is hard-coded in the app.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

KICK = [
    ("trsniIs", "Germany"), ("tittoshow", "Spain"), ("sussulevrai", "France"), ("nikkilve", "Germany"),
    ("teufeurs", "France"), ("kms", "France"), ("radiostafac", "Germany"), ("lemonmd", "France"),
    ("julktrainer", "Spain"), ("fekah", "France"), ("quentin-cey", "France"), ("balti", "France"),
    ("davidsantos-", "Spain"), ("botkz", "France"), ("kal-75020yt", "France"), ("just-ptit-flo", "France"),
    ("supersirc", "France"), ("ayzoh", "France"),
]
TWITCH = [
    ("bobinice7676", "France"), ("kosstochka", "Italy"), ("primeleague", "Germany"), ("starwraith", "Spain"),
    ("juakynen", "Spain"), ("nyaneila", "Germany"), ("lolitoIfdez", "Spain"), ("fierik", "Italy"),
    ("valleague_it", "Italy"), ("emerymoonvr", "Italy"), ("esportmaniacos", "Spain"), ("noni", "Spain"),
    ("marcomerrino", "Italy"), ("areliann", "France"), ("jimmyboyyy", "France"), ("wankilstudio", "France"),
    ("elbrokeron", "Spain"), ("fakemonster", "France"), ("creepy_live", "Italy"), ("iohuntuniverse", "Italy"),
]


def write(path: Path, headers: list[str], rows: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
    for source, country in rows:
        ws.append([source, country, None, None])
    for col, width in zip("ABCD", (20, 12, 22, 26), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"wrote {path} ({len(rows)} rows)")


if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "samples"
    write(out / "kick_streamers.xlsx", ["id_kick", "country", "id_twitch", "remarks"], KICK)
    write(out / "twitch_streamers.xlsx", ["id_twitch", "country", "id_kick", "remarks"], TWITCH)
